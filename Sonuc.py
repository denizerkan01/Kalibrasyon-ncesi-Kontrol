import openpyxl
import os
from openpyxl.styles import Alignment
from openpyxl.styles import Font
import datetime

HEADERS = {
    "A": ("Device SN\\Test Adımları", 30),
    "B": ("Tarih - Zaman", 30),
    "C": ("Kalibrasyon Matris Kontrolü", 30),
    "D": ("Firmware Sürüm Kontrolü", 30),
    "E": ("Acc X Testi Değerleri", 30),
    "F": ("Acc Y Testi Değerleri", 30),
    "G": ("Acc Z Testi Değerleri", 30),
    "H": ("Acc Testi Sonucu", 30),
    "I": ("Gyro X Testi Değerleri", 30),
    "J": ("Gyro Y Testi Değerleri", 30),
    "K": ("Gyro Z Testi Değerleri", 30),
    "L": ("Gyro Testi Sonucu", 30),
    "M": ("Magn Norm Testi Değerleri", 45),
    "N": ("Magn Norm Testi Sonucu", 45),
    "O": ("UART Bağlantı Testi Sonucu", 45),
    "P": ("RS 232 Bağlantı Testi Sonucu", 45),
    "Q": ("Gyro Log Dosyası / Tekrar Eden Satırlar", 60),
    "R": ("Gyro Tekrar Testi Sonucu", 45),
    "S": ("UART / RS232 Genel Sonuç", 45),


}

Result_Values = {

    "E": ("-0.6 ile 0.6 Arası", 30),
    "F": ("-0.6 ile 0.6 Arası", 30),
    "G": ("-10.302 ile -9.302 Arası", 30),
    "I": ("-0.4 ile 0.4 Arası", 30),
    "J": ("-0.4 ile 0.4 Arası", 30),
    "K": ("-0.4 ile 0.4 Arası", 30),
    "M": ("0.37 ile 0.57 Arası", 30),

}

def _ensure_headers(sheet):

    sheet.row_dimensions[1].height = 30
    for col, (text, width) in HEADERS.items():
        sheet.column_dimensions[col].width = width
        cell = sheet[f"{col}1"]
        if cell.value != text:
            cell.value = text
            cell.font = Font(size=14, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

def _ensure_res_val(sheet):

    sheet.row_dimensions[2].height = 30
    for col, (text, width) in Result_Values.items():
        sheet.column_dimensions[col].width = width
        cell = sheet[f"{col}2"]
        if cell.value != text:
            cell.value = text
            cell.font = Font(size=14, bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center')

def find_row_by_sn(sheet, sn):
    row = 3
    last_match = None #yeni eklendi
    while True:
        cell_value = sheet.cell(row=row, column=1).value
        if cell_value is None:
            return last_match
        if cell_value == sn:
            #return row
            last_match = row
        row += 1

def write_result(sheet, sn, column, value, red=False):
    row = find_row_by_sn(sheet, sn)
    if row is None:
        #print(f"Could not write result — {sn} not found in column A")
        return
    cell = sheet.cell(row=row, column=column)
    cell.value = value
    if red:
        cell.font=Font(color="FF0000")

def Sonuc(pn, sn, base_path, results_Kalibrasyon_Kontrol, results_Acc_Acilis, results_Gyro_Acilis, results_Magn, result_uart, result_Gyro_El):
    folder_path = os.path.join(base_path, "Test Sonuclari")
    os.makedirs(folder_path, exist_ok=True)
    excel_folder_path = os.path.join(folder_path, "TestSonuclari.xlsx")

    if os.path.exists(excel_folder_path):
        wb = openpyxl.load_workbook(excel_folder_path)
        sheet = wb[pn] if pn in wb.sheetnames else wb.create_sheet(pn)
        if "Sheet1" in wb.sheetnames and "Sheet1" != pn:
            default_sheet = wb["Sheet1"]
            if default_sheet.max_row == 1 and default_sheet.max_column == 1 and default_sheet["A1"].value is None:
                del wb["Sheet1"]
    else:
        wb = openpyxl.Workbook()
        sheet = wb.active
        sheet.title = pn

    _ensure_headers(sheet)  
    _ensure_res_val(sheet)

    cell = sheet.cell(row=2, column=1)
    value = "Beklenen Sonuç"
    cell.value = value
    cell.font = Font(size=14, bold=True)
    cell.alignment = Alignment(horizontal='center', vertical='center')

    row = 3
    while True:
        cell_value = sheet.cell(row=row, column=1).value

        if cell_value is None:
            sheet.cell(row=row, column=1).value = sn
            break

        #if cell_value == sn:
            #break

        row += 1

    current_time = datetime.datetime.now()

    write_result(sheet, sn, 2, current_time)

    

    if results_Kalibrasyon_Kontrol["calib_ctrl_success"] == True:
        write_result(sheet, sn, 3, "Geçti") 
    else:
        write_result(sheet, sn, 3, "Kaldı", red=True)

    if results_Kalibrasyon_Kontrol["firmware_version_match"] == True:
        write_result(sheet, sn, 4, "Geçti") 
    else:
        write_result(sheet, sn, 4, "Kaldı", red=True)

    min_accX = results_Acc_Acilis["MinAccX"]
    max_accX = results_Acc_Acilis["MaxAccX"]

    min_accY = results_Acc_Acilis["MinAccY"]
    max_accY = results_Acc_Acilis["MaxAccY"]

    min_accZ = results_Acc_Acilis["MinAccZ"]
    max_accZ = results_Acc_Acilis["MaxAccZ"]

    write_result(sheet, sn, 5, f"{min_accX} / {max_accX}")
    write_result(sheet, sn, 6, f"{min_accY} / {max_accY}")
    write_result(sheet, sn, 7, f"{min_accZ} / {max_accZ}")

    if results_Acc_Acilis["AccResult"] == True:
        write_result(sheet, sn, 8, "Geçti") 
    else:
        write_result(sheet, sn, 8, "Kaldı", red=True)

    min_gyroX = results_Gyro_Acilis["MinGyroX"]
    max_gyroX = results_Gyro_Acilis["MaxGyroX"]
    
    min_gyroY = results_Gyro_Acilis["MinGyroY"]
    max_gyroY = results_Gyro_Acilis["MaxGyroY"]
    
    min_gyroZ = results_Gyro_Acilis["MinGyroZ"]
    max_gyroZ = results_Gyro_Acilis["MaxGyroZ"]

    write_result(sheet, sn, 9, f"{min_gyroX} / {max_gyroX}")
    write_result(sheet, sn, 10, f"{min_gyroY} / {max_gyroY}")
    write_result(sheet, sn, 11, f"{min_gyroZ} / {max_gyroZ}")

    if results_Gyro_Acilis["GyroResult"] == True:
        write_result(sheet, sn, 12, "Geçti") 
    else:
        write_result(sheet, sn, 12, "Kaldı", red=True)

    
    min_magnNorm = results_Magn["MinMagnNorm"]
    max_magnNorm = results_Magn["MaxMagnNorm"]

    write_result(sheet, sn, 13, f"{min_magnNorm} / {max_magnNorm}")

    if results_Magn["MagnResult"] == True:
        write_result(sheet, sn, 14, "Geçti") 
    else:
        write_result(sheet, sn, 14, "Kaldı", red=True)

    if result_uart["Conn_Uart"] == True:
        write_result(sheet, sn, 15, "Geçti") 
    else:
        write_result(sheet, sn, 15, "Kaldı", red=True)

    if result_uart["Conn_232"] == True:
        write_result(sheet, sn, 16, "Geçti") 
    else:
        write_result(sheet, sn, 16, "Kaldı", red=True)

    
    gyro_rows = result_Gyro_El.get("row", [])
    if isinstance(gyro_rows, (list, tuple)):
        gyro_rows_text = ", ".join(str(row) for row in gyro_rows) if gyro_rows else "Yok"
    else:
        gyro_rows_text = str(gyro_rows) if gyro_rows else "Yok"

    if result_Gyro_El.get("Gyro_Success") is True:
        log_file_name = result_Gyro_El.get("log_file", "")
        if log_file_name:
            gyro_rows_text = f"{log_file_name} | {gyro_rows_text}"

    write_result(sheet, sn, 17, gyro_rows_text)

    if result_Gyro_El["Gyro_Success"] == False:
        write_result(sheet, sn, 18, "Geçti") 
    else:
        write_result(sheet, sn, 18, "Kaldı", red=True)

    if result_uart["Conn_Result"] == True:
        write_result(sheet, sn, 19, "Geçti") 
    else:
        write_result(sheet, sn, 19, "Kaldı", red=True)

    wb.save(excel_folder_path)
